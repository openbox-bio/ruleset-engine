from openpyxl import load_workbook
import csv
from pathlib import Path


def excel_sheet_to_csv(
    excel_path: str,
    sheet_name: str,
    output_csv: str | None = None,
):
    """
    Read an Excel sheet and write it to a CSV file.

    Intelligent behavior:
    ---------------------
    • If ANY cell contains a comma, a TAB delimiter is used
    • Otherwise, a COMMA delimiter is used
    • TRUE/FALSE values are written as literal TRUE/FALSE
    • Empty cells are written as empty fields

    Parameters
    ----------
    excel_path : str
        Path to the Excel (.xlsx) file
    sheet_name : str
        Sheet name to export
    output_csv : str, optional
        Output CSV filename. If not provided, derived from Excel filename.
    """

    excel_path = Path(excel_path)

    if output_csv is None:
        output_csv = excel_path.with_suffix("").name + f"_{sheet_name}.csv"

    wb = load_workbook(excel_path, data_only=True)

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in {excel_path}")

    ws = wb[sheet_name]

    rows = []
    contains_comma = False

    for row in ws.iter_rows(values_only=True):
        processed_row = []
        for cell in row:
            if cell is None:
                value = ""
            elif isinstance(cell, bool):
                # Preserve TRUE / FALSE
                value = "TRUE" if cell else "FALSE"
            else:
                value = str(cell)

            if "," in value:
                contains_comma = True

            processed_row.append(value)

        rows.append(processed_row)

    delimiter = "\t" if contains_comma else ","

    with open(output_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=delimiter)
        writer.writerows(rows)

    print(f"Written: {output_csv}")
    print(f"Delimiter used: {'TAB' if delimiter == chr(9) else 'COMMA'}")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Safe Export of data from Excel sheet to CSV.")
    parser.add_argument("--excel_file", help="Path to Excel file (.xlsx)")
    parser.add_argument("--sheet_name", help="Sheet name")
    parser.add_argument("--output", help="Output CSV file name", default=None)

    args = parser.parse_args()

    excel_sheet_to_csv(
        excel_path=args.excel_file,
        sheet_name=args.sheet_name,
        output_csv=args.output,
    )
